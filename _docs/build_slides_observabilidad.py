"""Construye slides-observabilidad.pptx con la skill md-to-pptx (modo avanzado) y añade,
detrás de cada slide con captura, una o dos slides con la captura en grande.

Estilo de las capturas (definido a mano en PowerPoint sobre las slides 18 y 19):
- Slide de contenido: captura a la derecha (x 6,96", y 1,43", 5,95" de ancho, borde gris) y
  pie de foto debajo (Montserrat 14 pt). Si no cabe, se reduce la altura.
- Slide ampliada: layout «White» sin título; captura en la caja x 0,25", y 0,40",
  10,61 x 6,63" con borde gris; las muy anchas se centran en vertical.

Uso (desde _docs/):
    python3 build_slides_observabilidad.py
    python3 build_slides_observabilidad.py slides-observabilidad-pptx.md slides-observabilidad.pptx

La skill se busca en MD2PPTX_SKILL (carpeta scripts/ de md-to-pptx) o en la ruta por defecto.
"""
import os
import sys

SKILL = os.environ.get("MD2PPTX_SKILL") or "/Users/ignaciosanchez/Library/Application Support/Claude/local-agent-mode-sessions/skills-plugin/05ead52c-581b-4beb-b4f7-102e4af27382/2de13103-0621-49d5-8ba2-e247facac554/skills/md-to-pptx/scripts"
sys.path.insert(0, SKILL)
from planner import build_plan  # noqa: E402
import builder as B  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img", "observabilidad") + os.sep

# título de la slide original -> [(imagen, pie de foto, nota del ponente)]
AMPLIADAS = {
    "Ver: la vista de dirección": [
        ("01-resumen.png", "Resumen: KPIs del periodo, estado de cada agente y coste frente al cap",
         "Señalar la fila de KPIs (81 % de autonomía, 2,1 % de overrides) y la tarjeta del agente de Reglas, marcada como degradada por superar su cap."),
    ],
    "Ver: trazabilidad de extremo a extremo": [
        ("03-trazabilidad.png", "Traza de un mensaje: cascada de agentes, datos de la traza y override humano",
         "Señalar la cascada (tiempo, modelo y tokens por agente) y el recuadro del override: quién corrigió, cuándo y por qué."),
    ],
    "Entender: explicabilidad": [
        ("04-traza-explicada.png", "Ficha explicada: qué llegó, qué hizo cada agente y por qué se escaló",
         "Leer de arriba abajo: qué llegó, qué hizo cada agente, la decisión con su motivo y el guardrail G-05 que la provocó."),
    ],
    "Entender: Reasoning Replay y What-if": [
        ("05-reasoning-replay.png", "Razonamiento registrado por agente: entrada, pasos y salida",
         "Señalar que no es el pensamiento interno del modelo, sino el registro estructurado que devuelve cada agente."),
        ("05b-replay-diff.png", "Replay What-if: misma decisión con un modelo más pequeño",
         "Señalar el resultado: misma decisión y mismo ramo, −82 % de coste y −41 % de latencia."),
    ],
    "Limitar: identidad y permisos por agente": [
        ("06-ficha-agente.png", "Ficha del agente de Reglas: identidad, permisos e histórico de cambios",
         "Señalar las columnas «Puede» y «No puede», y debajo el histórico de autonomía con su motivo y aprobador."),
    ],
    "Limitar: autonomía progresiva": [
        ("13-autonomia-auditoria.png", "Registro de auditoría de la bajada de nivel L3 → L2",
         "Señalar el motivo (override del 4,1 %), quién lo aprobó, que es reversible y el motivo marcado como «disparó la bajada»."),
        ("13b-autonomia-evolucion.png", "Evolución del nivel de autonomía del agente de Reglas",
         "Recorrer el gráfico: subidas en verde y bajadas en rojo. La autonomía se gana y se pierde con datos."),
    ],
    "Limitar: Trust Score y supervisión adaptativa": [
        ("14-agente-comportamiento.png", "Comportamiento por periodo y variables que gobiernan al agente",
         "Señalar cómo sube la precisión y bajan escalados y overrides semana a semana; debajo, cada variable con su efecto."),
    ],
    "Limitar: guardrails": [
        ("15-guardrail-g02.png", "Ficha del guardrail de importe (G-02): disparos por día, canal y ramo",
         "Señalar los KPIs (118 disparos, 33 % de los escalados, 14 min de resolución humana) y la distribución por canal y ramo."),
    ],
    "Limitar: kill switch y resiliencia": [
        ("02-agentes-kill-switch.png", "Tarjetas de agente con estado, nivel, coste frente al cap e interruptor de pausa",
         "Señalar el interruptor de cada agente y el estado «Degradado» del agente de Reglas."),
    ],
    "Pagar: caps y presupuesto": [
        ("17-cap03-consumo.png", "Cap diario del agente de Reglas (CAP-03): al 114 % del límite",
         "Señalar el medidor al 114 % y la barra roja de hoy, por encima del límite; a la derecha, el guardrail G-07 que actuó."),
        ("20-coste-mensual.png", "Coste mensual acumulado frente al cap y proyección de cierre",
         "Señalar que llevamos el 59 % del cap mensual y la proyección de cierre (428 € frente a 500 €)."),
    ],
    "Pagar: de la alerta a la acción": [
        ("19-caps-superados.png", "Consumo de cada cap frente al 100 %: superados, en aviso y dentro",
         "Señalar los tres caps en rojo por encima de la línea del 100 % y el que está en aviso."),
        ("17b-cap03-acciones.png", "Acciones correctivas de CAP-03: automática, simular, aplicar o pedir aprobación",
         "Recorrer los tres tipos de acción y su ahorro estimado; el cambio de presupuesto requiere aprobación del comité."),
    ],
    "Pagar: el modelo adecuado para cada agente": [
        ("10-finops-modelos.png", "Comparativa de modelos: precio, consumo, coste, latencia y calidad",
         "Señalar que gpt-5 concentra el 80 % del coste y lo usa un solo agente: es el candidato a optimizar."),
    ],
    "Antes de cambiar, simular": [
        ("18-nuevo-cap.png", "Nuevo cap: vista previa del consumo frente al nuevo límite",
         "Señalar la vista previa: cuántas veces se habría superado en 14 días y el aviso de que ya existe CAP-01."),
        ("16-nuevo-guardrail.png", "Nuevo guardrail: vista previa e impacto estimado",
         "Señalar el impacto estimado: unos 47 disparos en 14 días y el 0,28 % de los mensajes pasaría a una persona."),
    ],
    "Auditar: libro de registro": [
        ("11-historico.png", "Histórico de gobierno: overrides, incidentes, caps, replays y cambios de nivel",
         "Señalar los filtros por tipo y que cada evento lleva agente y responsable, incluido «Sistema»."),
    ],
    "Cumplimiento por diseño": [
        ("12-etiqueta-normativa.png", "Etiqueta normativa: qué exige el artículo y si ya está en vigor",
         "Leer el texto de la etiqueta: AI Act art. 14, exigible a alto riesgo desde 2027; aquí se aplica como buena práctica."),
    ],
    "Ver: la portada del panel": [
        ("00-inicio.png", "Inicio: una ficha por pregunta, con dos cifras en vivo, su semáforo y su norma",
         "Recorrer los grupos de izquierda a derecha y señalar el semáforo rojo de Knowledge bases y FinOps: el panel dice dónde mirar primero."),
    ],
    "Entender: evidencias verificadas": [
        ("41-evidencias.png", "Ficha del mensaje: diez evidencias resaltadas y la regla A5 que lo manda a revisión",
         "Pasar el ratón por una evidencia: el título cambia a su tipo. Señalar la regla A5 en rojo: la frase «fractura de clavícula y dos costillas» es la que lo frena."),
    ],
    "Conocer: el conocimiento también se degrada": [
        ("31-knowledge.png", "Knowledge bases: inventario con la salud de cada base y su tendencia",
         "Señalar las tres degradadas y la crítica: la red de talleres lleva nueve días sin sincronizar."),
        ("32-kb-comparativa.png", "Comparativa de configuraciones de Condicionados Auto: la recomendada recupera el recall",
         "Señalar que el filtro de vigencia devuelve el recall a 0,88: 340 documentos de 2024 competían con los de 2026."),
    ],
    "Conocer: rúbricas que vigilan el conocimiento": [
        ("33-rubrica.png", "Editor de la rúbrica de Condicionados Auto: criterios con peso y preguntas de referencia",
         "Señalar que cada guardado crea una versión nueva y que la rúbrica se ejecuta en cada reindexado: si empeora, no se publica."),
    ],
    "Auditar: la prueba en cada decisión": [
        ("42-respuesta-cruda.png", "Respuesta cruda de MSG-A-12 con el bloque _gobernanza que añade la plataforma",
         "Señalar que el bloque lo añade la plataforma, no el modelo, y que cada parte lleva su norma: AI Act art. 12, RGPD art. 9 y 22."),
    ],
    "Auditar: termómetro de cumplimiento": [
        ("34-termometro.png", "Termómetro de cumplimiento: cobertura por marco y matriz norma → control",
         "Señalar los controles en ámbar: minimización, EIPD pendiente y aviso de IA en WhatsApp. Un 100 % no sería creíble."),
        ("35-inventario-pii.png", "Ficha de cumplimiento de una traza: datos personales detectados y su valor en la traza",
         "Señalar «es diabético»: dato de salud que no hacía falta para decidir; es el argumento para seudonimizar antes del modelo."),
    ],
}


# Capturas compuestas (dos capturas apiladas) de las slides de contenido con pareja
COMPUESTAS = {
    "Entender: Reasoning Replay y What-if": "c21-replay.png",
    "Limitar: autonomía progresiva": "c24-autonomia.png",
    "Pagar: caps y presupuesto": "c28-caps.png",
    "Pagar: de la alerta a la acción": "c29-acciones.png",
    "Antes de cambiar, simular": "c31-simular.png",
}

# Geometría (pulgadas) tomada de las slides 18 y 19 editadas a mano
CAP_X, CAP_Y, CAP_W = 6.96, 1.43, 5.95      # captura en la slide de contenido
PIE_X, PIE_W, PIE_GAP, PIE_PT = 6.94, 6.23, 0.17, 14
AMP_X, AMP_Y, AMP_W, AMP_H = 0.25, 0.40, 10.61, 6.63  # caja de la slide ampliada


def borde_gris(pic):
    """Borde gris como en las slides 18 y 19: color de fondo al 50 %."""
    spPr = pic._element.spPr
    ln = B.etree.SubElement(spPr, B.qn("a:ln"))
    fill = B.etree.SubElement(ln, B.qn("a:solidFill"))
    clr = B.etree.SubElement(fill, B.qn("a:schemeClr"), val="bg1")
    B.etree.SubElement(clr, B.qn("a:lumMod"), val="50000")


def encajar(path, w, h):
    with B.Image.open(path) as im:
        iw, ih = im.size
    r = min(w / iw, h / ih)
    return iw * r, ih * r


class BuilderAmpliadas(B.Builder):
    def build_content(self, s):
        super().build_content(s)
        extras = AMPLIADAS.pop(s["title"], [])
        if not extras:
            return
        self.ajustar_captura(self.prs.slides[-1], s["title"], extras)
        for img, _pie, nota in extras:
            slide = self.prs.slides.add_slide(self.layout_by_name("White"))
            self.add_slide_number(slide)
            self.notes(slide, nota)
            w, h = encajar(IMG + img, AMP_W, AMP_H)
            pic = B.add_picture(slide, IMG + img, AMP_X, AMP_Y + (AMP_H - h) / 2, w, h)
            borde_gris(pic)

    def ajustar_captura(self, slide, titulo, extras):
        """Recoloca la captura de la slide de contenido y le pone borde y pie de foto."""
        pics = [sh for sh in slide.shapes if sh.shape_type == 13]
        for p in pics:
            p._element.getparent().remove(p._element)
        img = COMPUESTAS.get(titulo, extras[0][0])
        pie = extras[0][1] if len(extras) == 1 or titulo not in COMPUESTAS else f"Arriba: {extras[0][1]}. Abajo: {extras[1][1]}."
        pie_h = B.n_lines(pie, PIE_W - 0.1, "Montserrat", PIE_PT) * PIE_PT * 1.2 / 72 + 0.06
        w, h = encajar(IMG + img, CAP_W, B.CONTENT_BOTTOM - CAP_Y - PIE_GAP - pie_h)
        pic = B.add_picture(slide, IMG + img, CAP_X, CAP_Y, w, h)
        borde_gris(pic)
        tb = B.textbox(slide, PIE_X, CAP_Y + h + PIE_GAP, PIE_W, pie_h)
        p = tb.text_frame.paragraphs[0]
        B.set_spacing(p, 1.2, 0)
        B.add_runs(p, pie, PIE_PT, None, "Montserrat")


if __name__ == "__main__":
    md = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "slides-observabilidad-pptx.md")
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "slides-observabilidad.pptx")
    plan = build_plan(open(md, encoding="utf-8").read(), "avanzada")
    b = BuilderAmpliadas(plan, os.path.dirname(os.path.abspath(md)))
    b.build(out)
    from pptx import Presentation
    print("diapositivas:", len(Presentation(out).slides), "| sin usar:", list(AMPLIADAS), "| avisos:", b.warnings, b.report)
